"""Bounded XLSX edits and PDF generation; never executes model-generated code."""
import json
import os
import math
import re
import secrets
import sqlite3
import time
import zipfile
from io import BytesIO
from pathlib import Path
from html import escape
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

MAX_BYTES=8*1024*1024

def validate_archive(data):
    if len(data)>MAX_BYTES:raise ValueError('A planilha excede 8 MB.')
    with zipfile.ZipFile(BytesIO(data)) as archive:
        entries=archive.infolist()
        if len(entries)>2000 or sum(x.file_size for x in entries)>30*1024*1024:
            raise ValueError('Planilha expandida grande demais.')
        if any('vbaProject' in x.filename for x in entries):raise ValueError('Macros não são permitidas.')

def workbook(data):
    validate_archive(data)
    book=load_workbook(BytesIO(data),data_only=False,keep_links=False)
    if len(book.worksheets)>20 or any(s.max_row>5000 or s.max_column>100 for s in book):
        book.close();raise ValueError('Limite: 20 abas, 5.000 linhas e 100 colunas por aba.')
    return book

def snapshot(data):
    book=workbook(data);lines=[];count=0;size=0;truncated=False
    try:
        for sheet in book:
            lines.append('[ABA '+sheet.title+']')
            for row in sheet:
                for cell in row:
                    if cell.value is None:continue
                    line=f'{cell.coordinate}: {str(cell.value)[:250]}'
                    if count>=1800 or size+len(line)>24000:
                        truncated=True;break
                    lines.append(line);size+=len(line);count+=1
                if truncated:break
            if truncated:break
        if truncated:lines.append('[DADOS PARCIAIS: peça células/aba específicas se faltarem informações.]')
        return '\n'.join(lines)
    finally:book.close()

def parse_plan(raw):
    text=raw.strip()
    if text.startswith('```'):text='\n'.join(text.splitlines()[1:-1])
    try:plan=json.loads(text)
    except ValueError:raise ValueError('O modelo não produziu um plano válido. Nenhum arquivo foi alterado.') from None
    if not isinstance(plan,dict):raise ValueError('Plano inválido.')
    clarification=plan.get('clarification','')
    if clarification:
        if not isinstance(clarification,str):raise ValueError('Esclarecimento inválido.')
        raise ValueError(clarification[:1500])
    edits=plan.get('edits')
    if not isinstance(edits,list) or not 1<=len(edits)<=200:raise ValueError('O plano deve conter de 1 a 200 células.')
    return edits

def apply_edits(data,edits):
    book=workbook(data);seen=set();preview=[]
    try:
        for edit in edits:
            if not isinstance(edit,dict):raise ValueError('Alteração inválida.')
            sheet=edit.get('sheet');address=edit.get('cell');value=edit.get('value')
            if sheet not in book.sheetnames or not isinstance(address,str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]{0,3}',address):
                raise ValueError('Aba ou célula inválida.')
            cell=book[sheet][address]
            if cell.row>5000 or cell.column>100 or isinstance(cell,MergedCell):raise ValueError('Use a célula inicial da região mesclada, dentro dos limites.')
            key=(sheet,address)
            if key in seen:raise ValueError('Célula repetida no plano.')
            seen.add(key)
            if cell.data_type=='f':raise ValueError('Fórmulas existentes não podem ser sobrescritas por este fluxo.')
            if not (value is None or isinstance(value,(str,int,float,bool))):raise ValueError('Valor de célula inválido.')
            if isinstance(value,float) and not math.isfinite(value):raise ValueError('Número inválido.')
            if isinstance(value,str) and (len(value)>2000 or value.lstrip().startswith(('=','+','-','@'))):raise ValueError('Texto de célula inválido ou interpretável como fórmula.')
            preview.append(f'{sheet}!{address}: {str(cell.value)[:180]} → {str(value)[:300]}')
            cell.value=value
        output=BytesIO();book.save(output)
        if len(output.getvalue())>MAX_BYTES:raise ValueError('Arquivo gerado grande demais.')
        return output.getvalue(),'\n'.join(preview)
    finally:book.close()

def document_lines(text):
    if not isinstance(text,str) or not text.strip() or len(text)>40000:
        raise ValueError('Texto do documento inválido ou grande demais.')
    for line in text.splitlines():
        line=line.strip()
        if not line:
            yield 'space',''
        elif line.startswith('#'):
            match=re.match(r'^(#{1,3})\s+(.+)$',line)
            yield ('heading'+str(len(match[1])),match[2]) if match else ('body',line)
        elif line.startswith(('- ','* ')):
            yield 'bullet',line[2:]
        else:
            yield 'body',line

def make_pdf(text):
    output=BytesIO();styles=getSampleStyleSheet();story=[]
    for kind,content in document_lines(text):
        if kind=='space':
            story.append(Spacer(1,8));continue
        style=styles[{'heading1':'Title','heading2':'Heading2','heading3':'Heading3','bullet':'BodyText','body':'BodyText'}[kind]]
        story.append(Paragraph(escape(content),style,bulletText='•' if kind=='bullet' else None))
        story.append(Spacer(1,5))
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('Helvetica',8)
        canvas.drawString(36,22,'THIGAS AI')
        canvas.drawRightString(doc.pagesize[0]-36,22,str(doc.page))
        canvas.restoreState()
    SimpleDocTemplate(output,title='Documento THIGAS AI',leftMargin=48,rightMargin=48,
                      topMargin=48,bottomMargin=42).build(story,onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()

def make_docx(text):
    from docx import Document
    from docx.shared import Pt
    document=Document()
    document.styles['Normal'].font.name='Calibri'
    document.styles['Normal'].font.size=Pt(11)
    for kind,content in document_lines(text):
        if kind.startswith('heading'):
            document.add_heading(content,level=int(kind[-1])-1)
        elif kind=='bullet':
            document.add_paragraph(content,style='List Bullet')
        else:
            document.add_paragraph(content)
    output=BytesIO();document.save(output)
    return output.getvalue()

class Store:
    def __init__(self,directory):self.directory=Path(directory)
    def db(self):
        self.directory.mkdir(parents=True,exist_ok=True,mode=0o700)
        os.chmod(self.directory,0o700)
        con=sqlite3.connect(self.directory/'documents.sqlite3',timeout=20)
        os.chmod(self.directory/'documents.sqlite3',0o600)
        con.execute('CREATE TABLE IF NOT EXISTS files (owner TEXT PRIMARY KEY, name TEXT, data BLOB, created REAL)')
        con.execute('CREATE TABLE IF NOT EXISTS plans (id TEXT PRIMARY KEY, owner TEXT, name TEXT, data BLOB, state TEXT, created REAL)')
        con.execute('DELETE FROM files WHERE created<?',(time.time()-86400,))
        con.execute('DELETE FROM plans WHERE created<?',(time.time()-86400,))
        return con
    def attach(self,owner,name,data):
        with self.db() as con:con.execute('INSERT OR REPLACE INTO files VALUES (?,?,?,?)',(owner,name,data,time.time()))
    def attached(self,owner):
        with self.db() as con:return con.execute('SELECT name,data FROM files WHERE owner=?',(owner,)).fetchone()
    def propose(self,owner,name,data):
        pid=secrets.token_hex(12)
        with self.db() as con:
            if con.execute('SELECT count(*) FROM plans WHERE owner=?',(owner,)).fetchone()[0]>=20:raise ValueError('Limite de 20 arquivos por dia. Aguarde a expiração.')
            con.execute('INSERT INTO plans VALUES (?,?,?,?,?,?)',(pid,owner,name,data,'pending',time.time()))
        return pid
    def decide(self,pid,owner,approve):
        with self.db() as con:
            changed=con.execute("UPDATE plans SET state=? WHERE id=? AND owner=? AND state='pending'",('approved' if approve else 'cancelled',pid,owner))
            if not changed.rowcount:raise ValueError('Revisão expirada, já decidida ou pertencente a outro usuário.')
            return con.execute('SELECT name,data FROM plans WHERE id=? AND owner=?',(pid,owner)).fetchone()
    def approved(self,pid,owner):
        with self.db() as con:return con.execute("SELECT name,data FROM plans WHERE id=? AND owner=? AND state='approved'",(pid,owner)).fetchone()
    def clear(self,owner):
        with self.db() as con:
            con.execute('DELETE FROM files WHERE owner=?',(owner,));con.execute('DELETE FROM plans WHERE owner=?',(owner,))
