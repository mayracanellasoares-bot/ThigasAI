import json
import tempfile
import unittest
from io import BytesIO
from unittest.mock import patch
from openpyxl import Workbook,load_workbook
from pypdf import PdfReader
from test_app import gateway
import document_agent as documents

def sample():
    book=Workbook();book.active.title='Semana 1';book.active['B2']='Antes';book.active['B2'].number_format='@';book.active['C2']='=1+1';book.active.merge_cells('D2:E2');out=BytesIO();book.save(out);book.close();return out.getvalue()

class Documents(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.store=documents.Store(self.tmp.name);self.addCleanup(self.tmp.cleanup)
    def test_edit_preserves_original_formulas_and_format(self):
        raw=sample();edited,preview=documents.apply_edits(raw,[{'sheet':'Semana 1','cell':'B2','value':'Python'}])
        original=load_workbook(BytesIO(raw));copy=load_workbook(BytesIO(edited));self.addCleanup(original.close);self.addCleanup(copy.close)
        self.assertEqual(original.active['B2'].value,'Antes');self.assertEqual(copy.active['B2'].value,'Python');self.assertEqual(copy.active['C2'].value,'=1+1');self.assertEqual(copy.active['B2'].number_format,'@');self.assertIn('Semana 1!B2',preview)
    def test_invalid_cells_formula_and_duplicate_rejected(self):
        raw=sample()
        for edits in [[{'sheet':'Semana 1','cell':'C2','value':'x'}],[{'sheet':'Semana 1','cell':'E2','value':'x'}],[{'sheet':'Outra','cell':'B2','value':'x'}],[{'sheet':'Semana 1','cell':'B2','value':'=HYPERLINK("url")'}],[{'sheet':'Semana 1','cell':'B2','value':'x'}]*2]:
            with self.subTest(edits=edits),self.assertRaises(ValueError):documents.apply_edits(raw,edits)
    def test_snapshot_contains_real_coordinates(self):
        text=documents.snapshot(sample());self.assertIn('[ABA Semana 1]',text);self.assertIn('B2: Antes',text)
    def test_bad_plan_or_clarification_rejected(self):
        for raw in ['Aqui está seu arquivo',json.dumps({'edits':[],'clarification':'Qual semana?'}),json.dumps({'edits':[]})]:
            with self.assertRaises(ValueError):documents.parse_plan(raw)
    def test_approval_belongs_to_owner_and_is_single_use(self):
        pid=self.store.propose('1:1','x.xlsx',sample())
        with self.assertRaises(ValueError):self.store.decide(pid,'2:2',True)
        self.assertIsNone(self.store.approved(pid,'1:1'));self.store.decide(pid,'1:1',True)
        with self.assertRaises(ValueError):self.store.decide(pid,'1:1',True)
        self.assertIsNotNone(self.store.approved(pid,'1:1'))
    def test_cancel_prevents_recovery(self):
        pid=self.store.propose('1:1','x.xlsx',sample());self.store.decide(pid,'1:1',False);self.assertIsNone(self.store.approved(pid,'1:1'))
    def test_pdf_contains_real_text(self):
        text=''.join(p.extract_text() for p in PdfReader(BytesIO(documents.make_pdf('Planejamento\nAula de Python <teste>'))).pages);self.assertIn('Aula de Python',text)
    def test_telegram_review_does_not_deliver_workbook(self):
        self.store.attach('1:1','agenda.xlsx',sample())
        with patch.object(gateway,'document_store',self.store),patch.object(gateway,'ask_maritaca',return_value=json.dumps({'edits':[{'sheet':'Semana 1','cell':'B2','value':'Python'}]})),patch.object(gateway,'telegram_send_message'),patch.object(gateway,'telegram_api') as api,patch.object(gateway,'telegram_send_document') as send:
            gateway.prepare_telegram_workbook('1:1',1,'Preencha B2 com Python');send.assert_not_called();self.assertIn('inline_keyboard',api.call_args.args[1]['reply_markup'])
    def test_documents_disabled_without_allowlist(self):
        with patch.dict(gateway.os.environ,{'THIGAS_TELEGRAM_ALLOWED_USERS':''}),patch.object(gateway,'telegram_send_message') as send,patch.object(gateway,'telegram_download_document') as download:
            gateway.process_document_message({'from':{'id':1},'chat':{'id':1,'type':'private'},'document':{'file_name':'a.xlsx'}},'');download.assert_not_called();self.assertIn('THIGAS_TELEGRAM_ALLOWED_USERS',send.call_args.args[1])
    def test_callback_delivers_only_once(self):
        pid=self.store.propose('1:1','agenda.xlsx',sample());update={'id':'query','from':{'id':1},'message':{'chat':{'id':1,'type':'private'}},'data':'doc:approve:'+pid}
        with patch.dict(gateway.os.environ,{'THIGAS_TELEGRAM_ALLOWED_USERS':'1'}),patch.object(gateway,'document_store',self.store),patch.object(gateway,'telegram_api'),patch.object(gateway,'telegram_send_message'),patch.object(gateway,'telegram_send_document',return_value=True) as send:
            gateway.process_document_callback(update);gateway.process_document_callback(update);send.assert_called_once()
if __name__=='__main__':unittest.main()
