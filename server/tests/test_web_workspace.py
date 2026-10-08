import os
import tempfile
import unittest
from io import BytesIO
from unittest.mock import Mock, patch

from flask import Flask
from docx import Document

import document_agent
from web_workspace import install


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        # Even if old credentials remain in Render, no social login or Neon connection is used.
        env = {
            'THIGAS_SESSION_SECRET': 'a' * 48,
            'DATABASE_URL': 'postgresql://example.invalid/old_users',
            'GOOGLE_CLIENT_ID': 'old-client',
            'GOOGLE_CLIENT_SECRET': 'old-secret',
            'GITHUB_CLIENT_ID': 'old-client',
            'GITHUB_CLIENT_SECRET': 'old-secret',
        }
        self.env = patch.dict(os.environ, env)
        self.env.start()
        self.addCleanup(self.env.stop)
        self.app = Flask(__name__, static_folder='../static')
        self.app.config['TESTING'] = True
        self.compose = Mock(return_value='# Documento\nConteúdo revisado.')
        self.store = document_agent.Store(self.tmp.name + '/files')

        class GatewayError(Exception):
            pass

        self.gateway = dict(
            document_store=self.store,
            compose_reviewed_document=self.compose,
            extract_document_text=Mock(return_value=('Fonte real', False)),
            ask_maritaca=Mock(),
            GatewayError=GatewayError,
        )
        install(self.app, self.gateway)
        self.client = self.app.test_client()
        self.session_data = self.start_session(self.client)

    def start_session(self, client):
        response = client.get('/workspace/api/session', base_url='https://test.example')
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json['anonymous'])
        self.assertTrue(response.json['csrf'])
        with client.session_transaction(base_url='https://test.example') as state:
            self.assertEqual(len(state['visitor']), 48)
            return {'csrf': state['csrf'], 'owner': 'web-anon:' + state['visitor']}

    def post(self, path, **kwargs):
        return self.client.post(
            path,
            base_url='https://test.example',
            headers={'X-CSRF-Token': self.session_data['csrf'],
                     'Origin': 'https://test.example'},
            **kwargs,
        )

    def test_no_login_and_no_social_routes(self):
        self.assertEqual(self.client.get('/auth/login/google').status_code, 404)
        self.assertEqual(self.client.get('/auth/login/github').status_code, 404)
        self.assertEqual(self.client.get('/auth/login/meta').status_code, 404)
        self.assertEqual(self.client.get('/auth/me').status_code, 404)
        response = self.client.get('/workspace')
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'Sem cadastro e sem login', response.data)
        self.assertNotIn(b'Entrar com Google', response.data)

    def test_anonymous_session_is_stable_and_secure(self):
        second = self.start_session(self.client)
        self.assertEqual(second, self.session_data)
        fresh = self.app.test_client()
        self.assertNotEqual(self.start_session(fresh)['owner'], self.session_data['owner'])
        self.assertEqual(fresh.post('/workspace/api/generate',
                                    base_url='https://test.example', json={}).status_code, 403)
        self.assertEqual(self.client.post('/workspace/api/clear',
                                          base_url='https://test.example').status_code, 403)
        self.assertEqual(self.client.post('/workspace/api/clear',
                                          base_url='https://test.example',
                                          headers={'X-CSRF-Token': self.session_data['csrf'],
                                                   'Origin': 'https://evil.example'}).status_code, 403)
        self.assertEqual(self.client.get('/workspace/api/session').headers['Cache-Control'], 'no-store')

    def test_generate_preview_approve_and_download_isolated(self):
        response = self.post('/workspace/api/generate',
                             json={'instruction': 'Crie uma aula', 'format': 'docx'})
        self.assertEqual(response.status_code, 200)
        pid = response.json['id']
        url = '/workspace/api/download/' + pid
        self.assertEqual(self.client.get(url, base_url='https://test.example').status_code, 404)
        self.assertEqual(self.post('/workspace/api/decision/' + pid,
                                   json={'approve': True}).status_code, 200)
        downloaded = self.client.get(url, base_url='https://test.example')
        self.assertEqual(downloaded.status_code, 200)
        self.assertIn('Conteúdo revisado.',
                      [p.text for p in Document(BytesIO(downloaded.data)).paragraphs])
        other = self.app.test_client()
        other_session = self.start_session(other)
        self.assertNotEqual(other_session['owner'], self.session_data['owner'])
        self.assertEqual(other.get(url, base_url='https://test.example').status_code, 404)
        self.assertEqual(other.post('/workspace/api/decision/' + pid,
                                    base_url='https://test.example',
                                    json={'approve': True},
                                    headers={'X-CSRF-Token': other_session['csrf']}).status_code, 409)

    def test_source_upload_and_clear(self):
        response = self.post('/workspace/api/upload',
                             data={'file': (BytesIO(b'fixture'), 'source.pdf')})
        self.assertEqual(response.status_code, 200)
        response = self.post('/workspace/api/generate',
                             json={'instruction': 'Resuma', 'format': 'pdf', 'use_attachment': True})
        self.assertEqual(response.status_code, 200)
        self.compose.assert_called_once_with('Resuma', 'Fonte real')
        self.assertEqual(self.post('/workspace/api/clear', json={}).status_code, 200)
        self.assertIsNone(self.store.attached(self.session_data['owner']))

    def test_cancel_invalid_payload_and_quota(self):
        owner = self.session_data['owner']
        pid = self.store.propose(owner, 'test.pdf', b'pdf')
        self.assertEqual(self.post('/workspace/api/decision/' + pid,
                                   json={'approve': 'yes'}).status_code, 400)
        self.assertEqual(self.post('/workspace/api/decision/' + pid,
                                   json={'approve': False}).status_code, 200)
        self.assertIsNone(self.store.approved(pid, owner))
        self.assertEqual(self.post('/workspace/api/generate', json=[]).status_code, 400)
        for _ in range(19):
            self.store.propose(owner, 'test.pdf', b'pdf')
        self.assertEqual(self.post('/workspace/api/generate',
                                   json={'instruction': 'Teste', 'format': 'pdf'}).status_code, 429)
        self.compose.assert_not_called()

    def test_xlsx_edit_and_expired_source(self):
        from openpyxl import Workbook, load_workbook
        book = Workbook()
        book.active.title = 'Aula'
        book.active['A1'] = 'Antes'
        output = BytesIO()
        book.save(output)
        book.close()
        self.assertEqual(self.post('/workspace/api/upload',
                                   data={'file': (BytesIO(output.getvalue()), 'aula.xlsx')}).status_code, 200)
        self.gateway['ask_maritaca'].return_value = (
            '{"edits":[{"sheet":"Aula","cell":"A1","value":"Depois"}]}'
        )
        result = self.post('/workspace/api/generate',
                           json={'instruction': 'Atualize A1', 'format': 'xlsx', 'use_attachment': True})
        self.assertEqual(result.status_code, 200)
        pid = result.json['id']
        self.post('/workspace/api/decision/' + pid, json={'approve': True})
        data = self.client.get('/workspace/api/download/' + pid,
                               base_url='https://test.example').data
        edited = load_workbook(BytesIO(data))
        self.assertEqual(edited.active['A1'].value, 'Depois')
        edited.close()
        self.store.clear(self.session_data['owner'])
        result = self.post('/workspace/api/generate',
                           json={'instruction': 'Use a fonte', 'format': 'docx', 'use_attachment': True})
        self.assertEqual(result.status_code, 400)

    def test_chat_does_not_require_auth(self):
        @self.app.post('/chat')
        def chat():
            return {'answer': 'ok'}
        response = self.app.test_client().post('/chat', json={'message': 'olá'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['answer'], 'ok')

    def test_missing_session_secret_only_disables_anonymous_document_sessions(self):
        with patch.dict(os.environ, {'THIGAS_SESSION_SECRET': ''}):
            app = Flask('no_secret_app')
            install(app, self.gateway)
            response = app.test_client().get('/workspace/api/session')
            self.assertEqual(response.status_code, 503)
            self.assertIn('THIGAS_SESSION_SECRET', response.json['error'])


if __name__ == '__main__':
    unittest.main()
