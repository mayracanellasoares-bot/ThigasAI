import sys
import unittest
from pathlib import Path
from io import BytesIO
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as gateway
from pptx import Presentation
from docx import Document
from pypdf import PdfReader

class ChatExports(unittest.TestCase):
    def setUp(self):
        self.client = gateway.app.test_client()

    def test_roundtrip_documents(self):
        text = '# Apresentação\n## Educação\n- Aprender com tecnologia\n- Preservar conteúdo'
        for kind in ('pptx', 'docx', 'pdf'):
            with self.subTest(kind=kind):
                response = self.client.post('/document/render', json={'format': kind, 'text': text})
                self.assertEqual(response.status_code, 200)
                self.assertIn('attachment;', response.headers['Content-Disposition'])
                data = BytesIO(response.data)
                if kind == 'pptx':
                    self.assertGreater(len(Presentation(data).slides), 0)
                    extracted, _ = gateway.extract_document_text('slides.pptx', response.data)
                    self.assertIn('Aprender', extracted)
                elif kind == 'docx':
                    self.assertIn('Aprender', '\n'.join(p.text for p in Document(data).paragraphs))
                else:
                    self.assertIn('Aprender', ''.join(p.extract_text() for p in PdfReader(data).pages))
                response.close()

    def test_invalid_payloads(self):
        for payload in ([], {'format': 'exe', 'text': 'x'}, {'format': 'pdf', 'text': 'x'*40001}, {'format':'pptx','text':''}):
            response = self.client.post('/document/render', json=payload)
            self.assertEqual(response.status_code, 400)

    def test_slides_do_not_drop_long_body(self):
        text = '# Teste\n' + 'conteúdo ' * 160 + 'FINAL_PRESERVADO'
        data = gateway.document_agent.make_pptx(text)
        deck = Presentation(BytesIO(data))
        extracted = ' '.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('FINAL_PRESERVADO', extracted)
        self.assertEqual(extracted.count('conteúdo'),160)
